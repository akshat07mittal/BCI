import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from generate_retrospective import sprint_data
import os

def generate_user_stories():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "User Stories (3 Sprints)"
    
    # Simple styling
    hdr_font = Font(bold=True, color="FFFFFF", size=11)
    hdr_fill = PatternFill(start_color="3B5998", end_color="3B5998", fill_type="solid")
    center = Alignment(horizontal="center", vertical="center")
    left = Alignment(horizontal="left", vertical="center")
    thick = Side(style="thin", color="000000")
    border = Border(left=thick, right=thick, top=thick, bottom=thick)
    
    columns = ["Sprint", "ID", "Assignee", "User Story", "Status", "Priority", "Story Pts"]
    
    ws.column_dimensions["A"].width = 15
    ws.column_dimensions["B"].width = 10
    ws.column_dimensions["C"].width = 18
    ws.column_dimensions["D"].width = 54
    ws.column_dimensions["E"].width = 12
    ws.column_dimensions["F"].width = 12
    ws.column_dimensions["G"].width = 12
    
    # Headers
    for c, hdr in enumerate(columns, 1):
        cell = ws.cell(row=1, column=c, value=hdr)
        cell.font = hdr_font
        cell.fill = hdr_fill
        cell.alignment = center
        cell.border = border
        
    cur_row = 2
    for sp in sprint_data:
        for items in sp["stories"]:
            # Create a row consisting of: Sprint Name + 6-tuple story elements
            uid, desc, status, prio, pts, assignee = items
            row_data = [sp["name"].split(" — ")[0], uid, assignee, desc, status, prio, pts]
            
            for c, val in enumerate(row_data, 1):
                cell = ws.cell(row=cur_row, column=c, value=val)
                cell.border = border
                cell.alignment = left if c == 4 else center
                
                if c == 5:
                    if val == "Done":
                        cell.font = Font(color="008000", bold=True)
                    elif val == "Partial":
                        cell.font = Font(color="D2691E", bold=True)
                    else:
                        cell.font = Font(color="FF0000", bold=True)
                        
            cur_row += 1

    out_path = os.path.join(os.path.dirname(os.path.dirname(__file__)), "documents", "User_Stories.xlsx")
    wb.save(out_path)
    print(f"Created: {out_path}")

if __name__ == "__main__":
    generate_user_stories()
